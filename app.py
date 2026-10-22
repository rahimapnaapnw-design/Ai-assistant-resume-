"""AI Resume ATS Checker.

Upload a resume (PDF, DOCX or TXT) and get an ATS score plus concrete
improvement suggestions, powered by Google Gemini Flash.
"""

import hashlib
import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# "gemini-flash-latest" always points to the newest Flash model. You can pin a
# specific model (e.g. "gemini-2.5-flash") with the GEMINI_MODEL secret/env var
# or from the sidebar.
DEFAULT_MODEL = "gemini-flash-latest"
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 30000
MAX_JD_CHARS = 10000
MIN_RESUME_CHARS = 150

SCORE_CATEGORIES = {
    "keywords": "Keywords & Relevance",
    "formatting": "Formatting & ATS Compatibility",
    "content": "Content & Impact",
    "structure": "Structure & Sections",
    "readability": "Readability & Language",
}


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep content in tables (skills grids, two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ----------------------------------------------------------------------------
# Gemini helpers
# ----------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{job_description}\n\"\"\"\n"
        if job_description
        else "JOB DESCRIPTION: (not provided - evaluate against general ATS best practices "
        "and infer the most likely target role from the resume)\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Evaluate the resume below as an ATS would, and as a hiring manager would. Be honest, strict and specific.
Never invent facts about the candidate. Base everything only on the resume text.

Score each category from 0 to 100:
- keywords: relevant hard/soft skills, tools, role keywords{" matched against the job description" if job_description else ""}
- formatting: ATS-friendly layout (standard headings, no tables/columns/graphics issues evident from the text, consistent dates, contact info present)
- content: quantified achievements, action verbs, impact over duties
- structure: presence and order of key sections (contact, summary, experience, education, skills, projects)
- readability: clarity, concision, grammar, consistent tense, length

Also give an overall_score from 0 to 100 (weighted holistic ATS score).

Return ONLY valid JSON with exactly this shape:
{{
  "overall_score": <int 0-100>,
  "category_scores": {{"keywords": <int>, "formatting": <int>, "content": <int>, "structure": <int>, "readability": <int>}},
  "summary": "<2-3 sentence overall assessment>",
  "detected_role": "<most likely target role>",
  "strengths": ["<short bullet>", ...],
  "weaknesses": ["<short bullet>", ...],
  "missing_keywords": ["<keyword or skill>", ...],
  "improvements": [
    {{
      "section": "<resume section, e.g. Experience>",
      "priority": "High" | "Medium" | "Low",
      "issue": "<what is wrong>",
      "suggestion": "<specific fix>",
      "example": "<optional rewritten example line, based only on the resume content; may be empty>"
    }}
  ]
}}

Give 3-6 strengths, 3-6 weaknesses, up to 15 missing keywords and 5-10 improvements, ordered by priority.

{jd_block}
RESUME:
\"\"\"
{resume_text}
\"\"\"
"""


def parse_json_response(raw: str) -> dict:
    """Parse JSON from a model response, tolerating code fences and stray text."""
    if not raw or not raw.strip():
        raise ValueError("The AI returned an empty response.")
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            return json.loads(text[start : end + 1])
        raise ValueError("Could not parse the AI response as JSON.")


def _clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if v is not None and str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Make sure the result has every field the UI needs, with safe types."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected AI response format.")
    raw_cats = data.get("category_scores") or {}
    cats = {k: _clamp_score(raw_cats.get(k)) for k in SCORE_CATEGORIES}
    overall = data.get("overall_score")
    if overall is None and any(cats.values()):
        overall = sum(cats.values()) / len(cats)

    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "section": str(item.get("section", "General")).strip() or "General",
                "priority": priority,
                "issue": str(item.get("issue", "")).strip(),
                "suggestion": str(item.get("suggestion", "")).strip(),
                "example": str(item.get("example", "") or "").strip(),
            }
        )

    return {
        "overall_score": _clamp_score(overall),
        "category_scores": cats,
        "summary": str(data.get("summary", "")).strip(),
        "detected_role": str(data.get("detected_role", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "Your Gemini API key looks invalid or lacks permission. Please check it."
    if "429" in low or "quota" in low or "resource_exhausted" in low:
        return "Gemini rate limit or quota reached. Please wait a minute and try again."
    if "404" in low or "not found" in low:
        return "The selected Gemini model was not found. Try another model name in the sidebar."
    return f"Something went wrong: {msg}"


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def get_secret(name: str, default: str = "") -> str:
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass
    return os.environ.get(name, default)


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    overall = result["overall_score"]
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{overall}/100", score_label(overall), delta_color="off")
        st.progress(overall / 100)
        if result["detected_role"]:
            st.caption(f"Detected target role: **{result['detected_role']}**")
    with right:
        if result["summary"]:
            st.markdown("#### Summary")
            st.write(result["summary"])

    st.markdown("#### Score breakdown")
    cols = st.columns(len(SCORE_CATEGORIES))
    for col, (key, label) in zip(cols, SCORE_CATEGORIES.items()):
        score = result["category_scores"][key]
        col.metric(label, f"{score}")
        col.progress(score / 100)

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Strengths")
        for s in result["strengths"] or ["No strengths reported."]:
            st.markdown(f"- {s}")
    with c2:
        st.markdown("#### Weaknesses")
        for w in result["weaknesses"] or ["No weaknesses reported."]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.markdown("#### Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.markdown("#### Recommended improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    order = {"High": 0, "Medium": 1, "Low": 2}
    items = sorted(result["improvements"], key=lambda i: order[i["priority"]])
    if not items:
        st.info("No specific improvements returned.")
    for i, item in enumerate(items):
        title = f"{icons[item['priority']]} {item['priority']} · {item['section']}"
        with st.expander(title, expanded=(i < 3)):
            if item["issue"]:
                st.markdown(f"**Issue:** {item['issue']}")
            if item["suggestion"]:
                st.markdown(f"**Fix:** {item['suggestion']}")
            if item["example"]:
                st.markdown("**Example:**")
                st.code(item["example"], language=None)

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# Main app
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.write(
        "Upload your resume and get an ATS score with practical suggestions to improve it. "
        "Optionally paste a job description for a tailored match."
    )

    with st.sidebar:
        st.header("Settings")
        api_key = get_secret("GEMINI_API_KEY")
        if api_key:
            st.success("Gemini API key loaded from secrets.")
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        model = st.text_input("Gemini model", value=get_secret("GEMINI_MODEL", DEFAULT_MODEL))
        st.caption("Your resume is sent to the Gemini API for analysis and is not stored by this app.")

    col_a, col_b = st.columns(2)
    with col_a:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_b:
        job_description = st.text_area(
            "Job description (optional)",
            height=150,
            placeholder="Paste the job description here for a tailored ATS match...",
        )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            st.stop()

        try:
            data = uploaded.getvalue()
            resume_text = extract_resume_text(uploaded.name, data)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "export it as a text-based PDF or DOCX. (Note: ATS systems can't read images either!)"
            )
            st.stop()

        resume_text = resume_text[:MAX_RESUME_CHARS]
        job_description = job_description.strip()[:MAX_JD_CHARS]

        cache_key = hashlib.sha256(
            f"{model}|{resume_text}|{job_description}".encode("utf-8")
        ).hexdigest()
        cache = st.session_state.setdefault("results_cache", {})

        if cache_key in cache:
            st.session_state["result"] = cache[cache_key]
        else:
            with st.spinner("Analyzing your resume with Gemini..."):
                try:
                    result = analyze_resume(api_key, model.strip(), resume_text, job_description)
                except Exception as exc:
                    st.error(friendly_error(exc))
                    st.stop()
            cache[cache_key] = result
            st.session_state["result"] = result

    if "result" in st.session_state:
        st.divider()
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
